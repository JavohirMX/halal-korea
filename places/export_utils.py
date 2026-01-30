"""
Export utilities for HalalPlace data in multiple formats.
"""

import csv
import json
import io
from datetime import datetime
from typing import List, Dict, Any, Optional
from django.http import HttpResponse, StreamingHttpResponse
from django.contrib.gis.geos import Point
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import HalalPlace, BusinessHours, TimeSlot


class PlaceExporter:
    """Export HalalPlace data to various formats."""

    def __init__(self, queryset=None, include_business_hours=True):
        """
        Initialize exporter.

        Args:
            queryset: QuerySet of HalalPlace (default: all approved places)
            include_business_hours: Whether to include business hours data
        """
        self.queryset = queryset or HalalPlace.objects.filter(status="approved")
        self.include_business_hours = include_business_hours

    def export_json(self, pretty=True) -> str:
        """Export places as JSON string."""
        data = self._get_export_data()
        if pretty:
            return json.dumps(
                data, indent=2, ensure_ascii=False, default=self._json_serializer
            )
        return json.dumps(data, ensure_ascii=False, default=self._json_serializer)

    def export_csv(self) -> str:
        """Export places as CSV string."""
        output = io.StringIO()
        writer = csv.writer(output)

        # Write header
        headers = self._get_csv_headers()
        writer.writerow(headers)

        # Write data
        for place_data in self._get_export_data():
            row = self._flatten_for_csv(place_data)
            writer.writerow(row)

        return output.getvalue()

    def export_excel(self) -> bytes:
        """Export places as Excel binary data."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Halal Places"

        # Style header
        header_fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True)

        # Write headers
        headers = self._get_csv_headers()
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Write data
        for row_idx, place_data in enumerate(self._get_export_data(), 2):
            row = self._flatten_for_csv(place_data)
            for col, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col, value=value)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def export_geojson(self) -> str:
        """Export places as GeoJSON FeatureCollection."""
        features = []

        for place in self.queryset.select_related("business_hours").prefetch_related(
            "business_hours__time_slots"
        ):
            if place.location:
                feature = {
                    "type": "Feature",
                    "geometry": {
                        "type": "Point",
                        "coordinates": [place.location.x, place.location.y],
                    },
                    "properties": self._get_place_dict(place),
                }
                features.append(feature)

        geojson = {
            "type": "FeatureCollection",
            "features": features,
            "metadata": {
                "generated_at": datetime.now().isoformat(),
                "count": len(features),
                "source": "Halal Korea Admin Export",
            },
        }

        return json.dumps(
            geojson, indent=2, ensure_ascii=False, default=self._json_serializer
        )

    def _get_export_data(self) -> List[Dict[str, Any]]:
        """Get all places as list of dictionaries."""
        data = []
        queryset = self.queryset.select_related("business_hours").prefetch_related(
            "business_hours__time_slots"
        )

        for place in queryset:
            place_dict = self._get_place_dict(place)
            data.append(place_dict)

        return data

    def _get_place_dict(self, place: HalalPlace) -> Dict[str, Any]:
        """Convert a HalalPlace instance to a dictionary."""
        data = {
            "id": place.id,
            "name": place.name,
            "description": place.description,
            "category": place.category,
            "category_display": place.get_category_display(),
            "status": place.status,
            "status_display": place.get_status_display(),
            "address": place.address,
            "phone_number": place.phone_number or "",
            "website": place.website or "",
            "google_map_link": place.google_map_link or "",
            "kakao_map_link": place.kakao_map_link or "",
            "naver_map_link": place.naver_map_link or "",
            "latitude": place.location.y if place.location else None,
            "longitude": place.location.x if place.location else None,
            "photo_urls": place.photo_urls or [],
            "photo_count": len(place.photo_urls) if place.photo_urls else 0,
            "cached_average_rating": float(place.cached_average_rating)
            if place.cached_average_rating
            else None,
            "cached_reviews_count": place.cached_reviews_count,
            "temporary_closure_until": place.temporary_closure_until.isoformat()
            if place.temporary_closure_until
            else None,
            "temporary_closure_reason": place.temporary_closure_reason or "",
            "submitted_by": place.submitted_by.username if place.submitted_by else None,
            "created_at": place.created_at.isoformat(),
            "updated_at": place.updated_at.isoformat(),
        }

        # Add business hours if requested
        if self.include_business_hours and hasattr(place, "business_hours"):
            bh = place.business_hours
            if bh:
                data["business_hours"] = {
                    "is_24_hours": bh.is_24_hours,
                    "notes": bh.notes or "",
                    "schedule": self._get_business_hours_schedule(bh),
                }

        return data

    def _get_business_hours_schedule(
        self, business_hours: BusinessHours
    ) -> Dict[str, Any]:
        """Get business hours schedule as dictionary."""
        schedule = {}

        for slot in business_hours.time_slots.all():
            day_name = slot.get_day_of_week_display()
            if day_name not in schedule:
                schedule[day_name] = []

            if slot.is_closed:
                schedule[day_name] = "closed"
            else:
                schedule[day_name].append(
                    {
                        "open": slot.open_time.strftime("%H:%M")
                        if slot.open_time
                        else None,
                        "close": slot.close_time.strftime("%H:%M")
                        if slot.close_time
                        else None,
                    }
                )

        return schedule

    def _get_csv_headers(self) -> List[str]:
        """Get CSV column headers."""
        return [
            "id",
            "name",
            "description",
            "category",
            "status",
            "address",
            "phone_number",
            "website",
            "google_map_link",
            "kakao_map_link",
            "naver_map_link",
            "latitude",
            "longitude",
            "photo_urls",
            "photo_count",
            "cached_average_rating",
            "cached_reviews_count",
            "temporary_closure_until",
            "temporary_closure_reason",
            "submitted_by",
            "created_at",
            "updated_at",
        ]

    def _flatten_for_csv(self, place_data: Dict[str, Any]) -> List[Any]:
        """Flatten place data for CSV format."""
        # Convert photo_urls list to semicolon-separated string
        photo_urls_str = (
            ";".join(place_data["photo_urls"]) if place_data["photo_urls"] else ""
        )

        return [
            place_data["id"],
            place_data["name"],
            place_data["description"],
            place_data["category"],
            place_data["status"],
            place_data["address"],
            place_data["phone_number"],
            place_data["website"],
            place_data["google_map_link"],
            place_data["kakao_map_link"],
            place_data["naver_map_link"],
            place_data["latitude"],
            place_data["longitude"],
            photo_urls_str,
            place_data["photo_count"],
            place_data["cached_average_rating"],
            place_data["cached_reviews_count"],
            place_data["temporary_closure_until"],
            place_data["temporary_closure_reason"],
            place_data["submitted_by"],
            place_data["created_at"],
            place_data["updated_at"],
        ]

    def _json_serializer(self, obj):
        """Custom JSON serializer for special types."""
        if isinstance(obj, datetime):
            return obj.isoformat()
        if isinstance(obj, Point):
            return {"latitude": obj.y, "longitude": obj.x}
        raise TypeError(f"Object of type {type(obj)} is not JSON serializable")

    def get_http_response(
        self, format_type: str, filename: Optional[str] = None
    ) -> HttpResponse:
        """Get HTTP response with exported data."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if format_type == "json":
            content = self.export_json()
            if not filename:
                filename = f"halal_places_{timestamp}.json"
            response = HttpResponse(content, content_type="application/json")

        elif format_type == "csv":
            content = self.export_csv()
            if not filename:
                filename = f"halal_places_{timestamp}.csv"
            response = HttpResponse(content, content_type="text/csv")

        elif format_type == "excel":
            content = self.export_excel()
            if not filename:
                filename = f"halal_places_{timestamp}.xlsx"
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        elif format_type == "geojson":
            content = self.export_geojson()
            if not filename:
                filename = f"halal_places_{timestamp}.geojson"
            response = HttpResponse(content, content_type="application/geo+json")

        else:
            raise ValueError(f"Unsupported format: {format_type}")

        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
