"""
Tests for places import/export functionality.
Tests export utilities, import utilities, and data management views.
"""

import json
import csv
import io
from datetime import datetime, date
from decimal import Decimal

from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.contrib.auth import get_user_model
from django.contrib.gis.geos import Point
from django.core.files.uploadedfile import SimpleUploadedFile

from .models import HalalPlace, BusinessHours, TimeSlot
from .export_utils import PlaceExporter
from .import_utils import PlaceImporter, ImportResult

User = get_user_model()


class ExportUtilsTests(TestCase):
    """Test all 4 export formats: JSON, CSV, Excel, GeoJSON."""

    def setUp(self):
        self.user = User.objects.create_user(
            username="testuser",
            email="test@example.com",
            password="testpass123",
            is_staff=True,
        )
        self.place1 = HalalPlace.objects.create(
            name="Test Restaurant",
            description="A test halal restaurant",
            category="restaurant",
            location=Point(126.978, 37.5665),  # Seoul coordinates
            address="123 Test Street, Seoul",
            phone_number="02-123-4567",
            website="https://test.com",
            google_map_link="https://maps.google.com/?q=37.5665,126.9780",
            photo_urls=[
                "https://example.com/photo1.jpg",
                "https://example.com/photo2.jpg",
            ],
            status="approved",
            submitted_by=self.user,
            cached_average_rating=Decimal("4.5"),
            cached_reviews_count=10,
        )
        self.place2 = HalalPlace.objects.create(
            name="Test Market",
            description="A test halal market",
            category="market",
            location=Point(127.0, 37.55),
            address="456 Market Street, Seoul",
            status="pending",
            submitted_by=self.user,
        )
        # Create business hours for place1
        self.business_hours = BusinessHours.objects.create(
            place=self.place1,
            is_24_hours=False,
            notes="Test hours",
        )
        TimeSlot.objects.create(
            business_hours=self.business_hours,
            day_of_week=0,  # Monday
            open_time="09:00",
            close_time="22:00",
        )
        TimeSlot.objects.create(
            business_hours=self.business_hours,
            day_of_week=1,  # Tuesday
            is_closed=True,
        )

    def test_export_json_format(self):
        """Test JSON export format and content."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        json_output = exporter.export_json()

        # Verify it's valid JSON
        data = json.loads(json_output)
        self.assertIsInstance(data, list)
        self.assertEqual(len(data), 2)

        # Check first place data
        place1_data = next(p for p in data if p["id"] == self.place1.id)
        self.assertEqual(place1_data["name"], "Test Restaurant")
        self.assertEqual(place1_data["category"], "restaurant")
        self.assertEqual(place1_data["status"], "approved")
        self.assertEqual(place1_data["address"], "123 Test Street, Seoul")
        self.assertEqual(place1_data["latitude"], 37.5665)
        self.assertEqual(place1_data["longitude"], 126.978)
        self.assertEqual(place1_data["photo_count"], 2)
        self.assertEqual(float(place1_data["cached_average_rating"]), 4.5)
        self.assertEqual(place1_data["cached_reviews_count"], 10)

    def test_export_json_without_business_hours(self):
        """Test JSON export with business_hours disabled."""
        exporter = PlaceExporter(HalalPlace.objects.all(), include_business_hours=False)
        json_output = exporter.export_json()
        data = json.loads(json_output)

        place1_data = next(p for p in data if p["id"] == self.place1.id)
        self.assertNotIn("business_hours", place1_data)

    def test_export_csv_format(self):
        """Test CSV export format and content."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        csv_output = exporter.export_csv()

        # Parse CSV
        reader = csv.DictReader(io.StringIO(csv_output))
        rows = list(reader)

        self.assertEqual(len(rows), 2)

        # Find place1 in CSV
        place1_row = next(r for r in rows if r["id"] == str(self.place1.id))
        self.assertEqual(place1_row["name"], "Test Restaurant")
        self.assertEqual(place1_row["category"], "restaurant")
        self.assertEqual(place1_row["status"], "approved")
        self.assertEqual(place1_row["address"], "123 Test Street, Seoul")
        self.assertEqual(place1_row["phone_number"], "02-123-4567")
        self.assertEqual(place1_row["website"], "https://test.com")
        self.assertIn("photo_urls", place1_row)

    def test_export_csv_photo_urls_semicolon_separated(self):
        """Test that photo URLs are semicolon-separated in CSV."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        csv_output = exporter.export_csv()

        reader = csv.DictReader(io.StringIO(csv_output))
        rows = list(reader)
        place1_row = next(r for r in rows if r["id"] == str(self.place1.id))

        # Photo URLs should be semicolon-separated
        self.assertIn(";", place1_row["photo_urls"])
        self.assertIn("https://example.com/photo1.jpg", place1_row["photo_urls"])
        self.assertIn("https://example.com/photo2.jpg", place1_row["photo_urls"])

    def test_export_excel_format(self):
        """Test Excel export format and content."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        excel_output = exporter.export_excel()

        # Verify output is bytes
        self.assertIsInstance(excel_output, bytes)
        self.assertGreater(len(excel_output), 0)

        # Try to read with openpyxl
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(excel_output))
        ws = wb.active

        # Check sheet name
        self.assertEqual(ws.title, "Halal Places")

        # Check header row
        headers = [cell.value for cell in ws[1]]
        self.assertIn("id", headers)
        self.assertIn("name", headers)
        self.assertIn("category", headers)
        self.assertIn("status", headers)

        # Check data rows exist
        self.assertGreater(ws.max_row, 1)

    def test_export_geojson_format(self):
        """Test GeoJSON export format and content."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        geojson_output = exporter.export_geojson()

        # Verify it's valid JSON
        data = json.loads(geojson_output)

        # Check GeoJSON structure
        self.assertEqual(data["type"], "FeatureCollection")
        self.assertIn("features", data)
        self.assertIn("metadata", data)
        self.assertEqual(data["metadata"]["source"], "Halal Korea Admin Export")
        self.assertIn("generated_at", data["metadata"])

        # Check features
        features = data["features"]
        self.assertEqual(len(features), 2)

        # Check first feature structure
        feature = next(f for f in features if f["properties"]["id"] == self.place1.id)
        self.assertEqual(feature["type"], "Feature")
        self.assertEqual(feature["geometry"]["type"], "Point")
        self.assertEqual(feature["geometry"]["coordinates"], [126.978, 37.5665])
        self.assertEqual(feature["properties"]["name"], "Test Restaurant")

    def test_export_geojson_only_places_with_location(self):
        """Test that GeoJSON only includes places with location data."""
        # Note: location is required in the model, so we can't create places without it
        # Instead, we'll test that GeoJSON only exports places with valid coordinates
        # Create a place with location
        HalalPlace.objects.create(
            name="With Location Place",
            description="Test place with location",
            category="restaurant",
            location=Point(127.1, 37.6),
            address="789 With Loc Street",
            status="approved",
        )

        exporter = PlaceExporter(HalalPlace.objects.all())
        geojson_output = exporter.export_geojson()
        data = json.loads(geojson_output)

        # Should include all places with location (3 total now)
        self.assertEqual(len(data["features"]), 3)

    def test_get_http_response_json(self):
        """Test HTTP response for JSON export."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        response = exporter.get_http_response("json", "test_export.json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertIn(
            'attachment; filename="test_export.json"', response["Content-Disposition"]
        )

    def test_get_http_response_csv(self):
        """Test HTTP response for CSV export."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        response = exporter.get_http_response("csv", "test_export.csv")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertIn(
            'attachment; filename="test_export.csv"', response["Content-Disposition"]
        )

    def test_get_http_response_excel(self):
        """Test HTTP response for Excel export."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        response = exporter.get_http_response("excel", "test_export.xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            'attachment; filename="test_export.xlsx"', response["Content-Disposition"]
        )

    def test_get_http_response_geojson(self):
        """Test HTTP response for GeoJSON export."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        response = exporter.get_http_response("geojson", "test_export.geojson")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/geo+json")
        self.assertIn(
            'attachment; filename="test_export.geojson"',
            response["Content-Disposition"],
        )

    def test_get_http_response_invalid_format(self):
        """Test that invalid format raises ValueError."""
        exporter = PlaceExporter(HalalPlace.objects.all())
        with self.assertRaises(ValueError):
            exporter.get_http_response("invalid_format")


class ImportUtilsTests(TestCase):
    """Test JSON import, CSV import, updates, validation, and dry-run."""

    def setUp(self):
        self.user = User.objects.create_user(
            username="testuser",
            email="test@example.com",
            password="testpass123",
        )
        self.existing_place = HalalPlace.objects.create(
            name="Existing Place",
            description="An existing place",
            category="restaurant",
            location=Point(126.978, 37.5665),
            address="123 Existing Street",
            status="approved",
            submitted_by=self.user,
        )

    def test_import_json_creates_new_place(self):
        """Test importing JSON creates a new place."""
        json_data = [
            {
                "name": "New JSON Place",
                "description": "A new place from JSON",
                "category": "market",
                "address": "456 New Street",
                "latitude": 37.55,
                "longitude": 127.0,
                "status": "pending",
            }
        ]

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(json_data))

        # Debug: print errors if any
        if result.error_count > 0:
            print(f"Import errors: {result.errors}")

        self.assertEqual(
            result.created_count,
            1,
            f"Expected 1 created, got {result.created_count}. Errors: {result.errors}",
        )
        self.assertEqual(result.updated_count, 0)
        self.assertEqual(
            result.error_count,
            0,
            f"Expected 0 errors, got {result.error_count}: {result.errors}",
        )

        # Verify place was created
        place = HalalPlace.objects.get(name="New JSON Place")
        self.assertEqual(place.category, "market")
        self.assertEqual(place.address, "456 New Street")
        self.assertEqual(place.status, "pending")
        self.assertEqual(place.submitted_by, self.user)

    def test_import_json_updates_existing_place(self):
        """Test importing JSON updates an existing place."""
        json_data = [
            {
                "id": self.existing_place.id,
                "name": "Updated Name",
                "description": "Updated description",
                "category": "mosque",
                "address": "123 Existing Street",  # Same address
                "latitude": 37.5665,
                "longitude": 126.978,
                "status": "approved",
            }
        ]

        importer = PlaceImporter(user=self.user, update_existing=True)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.updated_count, 1)

        # Verify place was updated
        self.existing_place.refresh_from_db()
        self.assertEqual(self.existing_place.name, "Updated Name")
        self.assertEqual(self.existing_place.description, "Updated description")
        self.assertEqual(self.existing_place.category, "mosque")

    def test_import_json_skips_existing_when_update_false(self):
        """Test that existing places are skipped when update_existing=False."""
        json_data = [
            {
                "id": self.existing_place.id,
                "name": "Should Not Update",
                "description": "Should not be saved",
                "category": "market",
                "address": "123 Existing Street",
                "latitude": 37.5665,
                "longitude": 126.978,
                "status": "approved",
            }
        ]

        importer = PlaceImporter(user=self.user, update_existing=False)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.skipped_count, 1)
        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.updated_count, 0)

        # Verify place was NOT updated
        self.existing_place.refresh_from_db()
        self.assertEqual(self.existing_place.name, "Existing Place")

    def test_import_json_dry_run(self):
        """Test dry-run mode doesn't save to database."""
        initial_count = HalalPlace.objects.count()

        json_data = [
            {
                "name": "Dry Run Place",
                "description": "Should not be saved",
                "category": "restaurant",
                "address": "789 Dry Run Street",
                "latitude": 37.5,
                "longitude": 127.5,
                "status": "pending",
            }
        ]

        importer = PlaceImporter(user=self.user, dry_run=True)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.created_count, 1)
        # Verify no new place in database
        self.assertEqual(HalalPlace.objects.count(), initial_count)
        with self.assertRaises(HalalPlace.DoesNotExist):
            HalalPlace.objects.get(name="Dry Run Place")

    def test_import_json_validates_required_fields(self):
        """Test that required fields are validated."""
        json_data = [
            {
                # Missing name
                "description": "Missing name",
                "category": "restaurant",
                "address": "123 Test Street",
            }
        ]

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.error_count, 1)
        self.assertEqual(len(result.errors), 1)
        self.assertEqual(result.errors[0]["field"], "name")

    def test_import_json_validates_category(self):
        """Test that invalid category is rejected."""
        json_data = [
            {
                "name": "Test Place",
                "description": "Test",
                "category": "invalid_category",
                "address": "123 Test Street",
            }
        ]

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.error_count, 1)
        self.assertEqual(result.errors[0]["field"], "category")

    def test_import_json_validates_status(self):
        """Test that invalid status is rejected."""
        json_data = [
            {
                "name": "Test Place",
                "description": "Test",
                "category": "restaurant",
                "address": "123 Test Street",
                "status": "invalid_status",
            }
        ]

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.error_count, 1)
        self.assertEqual(result.errors[0]["field"], "status")

    def test_import_csv_creates_new_place(self):
        """Test importing CSV creates new places."""
        csv_content = """name,description,category,address,latitude,longitude,status,phone_number
New CSV Place,From CSV,market,456 CSV Street,37.55,127.0,pending,02-987-6543"""

        importer = PlaceImporter(user=self.user)
        result = importer.import_csv(csv_content)

        self.assertEqual(result.created_count, 1)
        self.assertEqual(result.error_count, 0)

        # Verify place was created
        place = HalalPlace.objects.get(name="New CSV Place")
        self.assertEqual(place.category, "market")
        self.assertEqual(place.phone_number, "02-987-6543")

    def test_import_csv_handles_photo_urls(self):
        """Test CSV import with semicolon-separated photo URLs."""
        csv_content = """name,description,category,address,latitude,longitude,status,photo_urls
Photo Place,With photos,restaurant,123 Photo St,37.5,127.5,pending,https://a.jpg;https://b.jpg;https://c.jpg"""

        importer = PlaceImporter(user=self.user)
        result = importer.import_csv(csv_content)

        self.assertEqual(result.created_count, 1)

        place = HalalPlace.objects.get(name="Photo Place")
        self.assertEqual(len(place.photo_urls), 3)
        self.assertIn("https://a.jpg", place.photo_urls)

    def test_import_csv_handles_dates(self):
        """Test CSV import with date fields."""
        csv_content = """name,description,category,address,latitude,longitude,status,temporary_closure_until,temporary_closure_reason
Closed Place,Temporarily closed,restaurant,123 Closed St,37.5,127.5,pending,2024-12-25,Holiday renovation"""

        importer = PlaceImporter(user=self.user)
        result = importer.import_csv(csv_content)

        self.assertEqual(result.created_count, 1)

        place = HalalPlace.objects.get(name="Closed Place")
        self.assertEqual(place.temporary_closure_until, date(2024, 12, 25))
        self.assertEqual(place.temporary_closure_reason, "Holiday renovation")

    def test_import_csv_handles_coordinates(self):
        """Test CSV import with coordinate conversion to Point."""
        csv_content = """name,description,category,address,latitude,longitude,status
Coord Place,With coordinates,restaurant,123 Coord St,37.1234,127.5678,pending"""

        importer = PlaceImporter(user=self.user)
        result = importer.import_csv(csv_content)

        self.assertEqual(result.created_count, 1)

        place = HalalPlace.objects.get(name="Coord Place")
        self.assertIsNotNone(place.location)
        self.assertAlmostEqual(place.location.y, 37.1234, places=4)
        self.assertAlmostEqual(place.location.x, 127.5678, places=4)

    def test_import_csv_skips_invalid_coordinates(self):
        """Test CSV import handles invalid coordinates gracefully by reporting error."""
        csv_content = """name,description,category,address,latitude,longitude,status
Invalid Coord,Invalid coordinates,restaurant,123 Invalid St,invalid,127.5,pending"""

        importer = PlaceImporter(user=self.user)
        result = importer.import_csv(csv_content)

        # Should report error for invalid coordinates, not create the place
        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.error_count, 1)
        self.assertEqual(result.errors[0]["field"], "csv_parse")

    def test_import_result_class(self):
        """Test ImportResult helper class."""
        result = ImportResult()

        result.add_error(1, "name", "Name is required", {"name": ""})
        result.add_error(2, "category", "Invalid category", {"category": "foo"})
        result.add_warning("Some warning")

        self.assertEqual(result.error_count, 2)
        self.assertEqual(len(result.errors), 2)
        self.assertEqual(len(result.warnings), 1)

        result_dict = result.to_dict()
        self.assertEqual(result_dict["errors"], 2)
        self.assertEqual(result_dict["created"], 0)
        self.assertEqual(len(result_dict["error_details"]), 2)

    def test_validate_file_json_valid(self):
        """Test validate_file with valid JSON."""
        json_data = [
            {
                "name": "Valid Place",
                "description": "Test",
                "category": "restaurant",
                "address": "123 Test St",
            }
        ]

        importer = PlaceImporter()
        is_valid, errors = importer.validate_file(json.dumps(json_data), "json")

        self.assertTrue(is_valid)
        self.assertEqual(len(errors), 0)

    def test_validate_file_json_invalid(self):
        """Test validate_file with invalid JSON."""
        importer = PlaceImporter()
        is_valid, errors = importer.validate_file("not valid json", "json")

        self.assertFalse(is_valid)
        self.assertGreater(len(errors), 0)

    def test_validate_file_csv_valid(self):
        """Test validate_file with valid CSV."""
        csv_content = """name,description,category,address
Valid CSV,Test,restaurant,123 Test St"""

        importer = PlaceImporter()
        is_valid, errors = importer.validate_file(csv_content, "csv")

        self.assertTrue(is_valid)
        self.assertEqual(len(errors), 0)

    def test_validate_file_csv_empty(self):
        """Test validate_file with empty CSV."""
        csv_content = """name,description,category,address"""

        importer = PlaceImporter()
        is_valid, errors = importer.validate_file(csv_content, "csv")

        self.assertFalse(is_valid)
        self.assertIn("CSV file is empty", errors[0])

    def test_validate_file_missing_required_fields(self):
        """Test validate_file detects missing required fields."""
        json_data = [{"description": "Missing name and address"}]

        importer = PlaceImporter()
        is_valid, errors = importer.validate_file(json.dumps(json_data), "json")

        self.assertFalse(is_valid)
        self.assertTrue(any("name" in e for e in errors))
        self.assertTrue(any("address" in e for e in errors))

    def test_find_existing_place_by_id(self):
        """Test finding existing place by ID."""
        importer = PlaceImporter()
        existing = importer._find_existing_place({"id": self.existing_place.id})

        self.assertEqual(existing.id, self.existing_place.id)

    def test_find_existing_place_by_name_address(self):
        """Test finding existing place by name and address."""
        importer = PlaceImporter()
        existing = importer._find_existing_place(
            {
                "name": "Existing Place",
                "address": "123 Existing Street",
            }
        )

        self.assertEqual(existing.id, self.existing_place.id)

    def test_find_existing_place_case_insensitive(self):
        """Test finding existing place with case-insensitive match."""
        importer = PlaceImporter()
        existing = importer._find_existing_place(
            {
                "name": "EXISTING PLACE",  # Different case
                "address": "123 EXISTING STREET",  # Different case
            }
        )

        self.assertEqual(existing.id, self.existing_place.id)


class DataManagementViewsTests(TestCase):
    """Test all data management views with proper permissions."""

    def setUp(self):
        self.client = Client()
        self.staff_user = User.objects.create_user(
            username="staffuser",
            email="staff@example.com",
            password="staffpass123",
            is_staff=True,
        )
        self.regular_user = User.objects.create_user(
            username="regularuser",
            email="regular@example.com",
            password="regularpass123",
            is_staff=False,
        )
        self.place = HalalPlace.objects.create(
            name="Test Place",
            description="A test place",
            category="restaurant",
            location=Point(126.978, 37.5665),
            address="123 Test Street",
            status="approved",
        )

    def test_dashboard_requires_staff(self):
        """Test that dashboard requires staff access."""
        # Anonymous user
        response = self.client.get(reverse("places:data_management_dashboard"))
        self.assertEqual(response.status_code, 302)  # Redirect to login

        # Regular user
        self.client.login(username="regularuser", password="regularpass123")
        response = self.client.get(reverse("places:data_management_dashboard"))
        self.assertEqual(response.status_code, 302)  # Redirect

        # Staff user
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:data_management_dashboard"))
        self.assertEqual(response.status_code, 200)

    def test_dashboard_shows_statistics(self):
        """Test dashboard shows place statistics."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:data_management_dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Test Place")
        # Should show counts
        self.assertContains(response, "total_places")

    def test_export_view_get(self):
        """Test export view GET request."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:export_places"))

        self.assertEqual(response.status_code, 200)
        # Should contain format choices
        self.assertContains(response, "JSON")
        self.assertContains(response, "CSV")
        self.assertContains(response, "Excel")
        self.assertContains(response, "GeoJSON")

    def test_export_view_post_json(self):
        """Test export view POST with JSON format."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(
            reverse("places:export_places"),
            {"format": "json", "status": "all", "include_business_hours": "on"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertIn("Content-Disposition", response)

    def test_export_view_post_csv(self):
        """Test export view POST with CSV format."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(
            reverse("places:export_places"),
            {"format": "csv", "status": "all"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_export_view_post_excel(self):
        """Test export view POST with Excel format."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(
            reverse("places:export_places"),
            {"format": "excel", "status": "all"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_export_view_post_geojson(self):
        """Test export view POST with GeoJSON format."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(
            reverse("places:export_places"),
            {"format": "geojson", "status": "all"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/geo+json")

    def test_export_view_invalid_format(self):
        """Test export view with invalid format shows error."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(
            reverse("places:export_places"),
            {"format": "invalid", "status": "all"},
        )

        # Should redirect with error message
        self.assertEqual(response.status_code, 302)

    def test_import_view_get(self):
        """Test import view GET request."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:import_places"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Import")

    def test_import_view_post_json(self):
        """Test import view POST with JSON file."""
        self.client.login(username="staffuser", password="staffpass123")

        json_data = json.dumps(
            [
                {
                    "name": "Imported JSON Place",
                    "description": "From import",
                    "category": "restaurant",
                    "address": "123 Import Street",
                    "latitude": 37.5,
                    "longitude": 127.5,
                    "status": "pending",
                }
            ]
        )

        uploaded_file = SimpleUploadedFile(
            "places.json", json_data.encode("utf-8"), content_type="application/json"
        )

        response = self.client.post(
            reverse("places:import_places"),
            {
                "import_file": uploaded_file,
                "format": "json",
                "dry_run": "off",
                "update_existing": "on",
            },
        )

        # Should redirect to dashboard
        self.assertEqual(response.status_code, 302)

        # Verify place was created
        place = HalalPlace.objects.get(name="Imported JSON Place")
        self.assertEqual(place.address, "123 Import Street")

    def test_import_view_post_csv(self):
        """Test import view POST with CSV file."""
        self.client.login(username="staffuser", password="staffpass123")

        csv_content = """name,description,category,address,latitude,longitude,status
Imported CSV Place,From CSV import,market,456 CSV Import,37.6,127.6,pending"""

        uploaded_file = SimpleUploadedFile(
            "places.csv", csv_content.encode("utf-8"), content_type="text/csv"
        )

        response = self.client.post(
            reverse("places:import_places"),
            {
                "import_file": uploaded_file,
                "format": "csv",
                "dry_run": "off",
                "update_existing": "on",
            },
        )

        self.assertEqual(response.status_code, 302)

        # Verify place was created
        place = HalalPlace.objects.get(name="Imported CSV Place")
        self.assertEqual(place.category, "market")

    def test_import_view_no_file(self):
        """Test import view with no file shows error."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(reverse("places:import_places"), {})

        # Should redirect with error
        self.assertEqual(response.status_code, 302)

    def test_import_view_dry_run(self):
        """Test import view with dry-run option."""
        self.client.login(username="staffuser", password="staffpass123")

        initial_count = HalalPlace.objects.count()

        json_data = json.dumps(
            [
                {
                    "name": "Dry Run Place",
                    "description": "Should not be saved",
                    "category": "restaurant",
                    "address": "789 Dry Run Street",
                    "status": "pending",
                }
            ]
        )

        uploaded_file = SimpleUploadedFile(
            "places.json", json_data.encode("utf-8"), content_type="application/json"
        )

        response = self.client.post(
            reverse("places:import_places"),
            {
                "import_file": uploaded_file,
                "format": "json",
                "dry_run": "on",
                "update_existing": "on",
            },
        )

        self.assertEqual(response.status_code, 302)
        # Count should not have changed
        self.assertEqual(HalalPlace.objects.count(), initial_count)

    def test_validate_ajax_endpoint_valid(self):
        """Test AJAX validation endpoint with valid file."""
        self.client.login(username="staffuser", password="staffpass123")

        json_data = json.dumps(
            [
                {
                    "name": "Valid Place",
                    "description": "Test",
                    "category": "restaurant",
                    "address": "123 Test St",
                }
            ]
        )

        uploaded_file = SimpleUploadedFile(
            "places.json", json_data.encode("utf-8"), content_type="application/json"
        )

        response = self.client.post(
            reverse("places:validate_import_file"),
            {"file": uploaded_file, "format": "json"},
        )

        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)
        self.assertTrue(data["valid"])
        self.assertEqual(len(data["errors"]), 0)

    def test_validate_ajax_endpoint_invalid(self):
        """Test AJAX validation endpoint with invalid file."""
        self.client.login(username="staffuser", password="staffpass123")

        uploaded_file = SimpleUploadedFile(
            "places.json", b"not valid json", content_type="application/json"
        )

        response = self.client.post(
            reverse("places:validate_import_file"),
            {"file": uploaded_file, "format": "json"},
        )

        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)
        self.assertFalse(data["valid"])
        self.assertGreater(len(data["errors"]), 0)

    def test_validate_ajax_no_file(self):
        """Test AJAX validation endpoint with no file."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.post(reverse("places:validate_import_file"), {})

        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)
        self.assertFalse(data["valid"])

    def test_import_template_download_csv(self):
        """Test CSV template download."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:import_template", args=["csv"]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertIn("Content-Disposition", response)

        # Verify content
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreater(len(rows), 0)
        self.assertIn("name", rows[0])

    def test_import_template_download_json(self):
        """Test JSON template download."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:import_template", args=["json"]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertIn("Content-Disposition", response)

        # Verify content
        data = json.loads(response.content)
        self.assertIn("places", data)
        self.assertGreater(len(data["places"]), 0)

    def test_import_template_invalid_format(self):
        """Test template download with invalid format."""
        self.client.login(username="staffuser", password="staffpass123")
        response = self.client.get(reverse("places:import_template", args=["invalid"]))

        # Should redirect with error
        self.assertEqual(response.status_code, 302)

    def test_non_staff_cannot_access_export(self):
        """Test that non-staff users cannot access export view."""
        self.client.login(username="regularuser", password="regularpass123")
        response = self.client.get(reverse("places:export_places"))
        self.assertEqual(response.status_code, 302)

    def test_non_staff_cannot_access_import(self):
        """Test that non-staff users cannot access import view."""
        self.client.login(username="regularuser", password="regularpass123")
        response = self.client.get(reverse("places:import_places"))
        self.assertEqual(response.status_code, 302)


class GeoJSONImportTests(TestCase):
    """Test GeoJSON FeatureCollection parsing."""

    def setUp(self):
        self.user = User.objects.create_user(
            username="testuser",
            email="test@example.com",
            password="testpass123",
        )
        self.existing_place = HalalPlace.objects.create(
            name="Existing Geo Place",
            description="An existing place",
            category="restaurant",
            location=Point(126.978, 37.5665),
            address="123 Existing Street",
            status="approved",
            submitted_by=self.user,
        )

    def test_import_geojson_featurecollection(self):
        """Test importing from GeoJSON FeatureCollection format."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    "geometry": {
                        "type": "Point",
                        "coordinates": [127.0, 37.55],
                    },
                    "properties": {
                        "name": "GeoJSON Place",
                        "description": "From GeoJSON",
                        "category": "market",
                        "address": "456 Geo Street",
                        "status": "pending",
                    },
                }
            ],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        self.assertEqual(result.created_count, 1)
        self.assertEqual(result.error_count, 0)

        # Verify place was created with coordinates
        place = HalalPlace.objects.get(name="GeoJSON Place")
        self.assertIsNotNone(place.location)
        self.assertAlmostEqual(place.location.x, 127.0, places=4)
        self.assertAlmostEqual(place.location.y, 37.55, places=4)

    def test_import_geojson_multiple_features(self):
        """Test importing multiple features from GeoJSON."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    "geometry": {"type": "Point", "coordinates": [127.0, 37.55]},
                    "properties": {
                        "name": "Place 1",
                        "description": "First place",
                        "category": "restaurant",
                        "address": "111 First St",
                        "status": "pending",
                    },
                },
                {
                    "type": "Feature",
                    "geometry": {"type": "Point", "coordinates": [127.1, 37.56]},
                    "properties": {
                        "name": "Place 2",
                        "description": "Second place",
                        "category": "mosque",
                        "address": "222 Second St",
                        "status": "approved",
                    },
                },
            ],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        self.assertEqual(result.created_count, 2)

    def test_import_geojson_updates_existing_by_id(self):
        """Test GeoJSON import updates existing place by ID."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    "geometry": {"type": "Point", "coordinates": [127.0, 37.55]},
                    "properties": {
                        "id": self.existing_place.id,
                        "name": "Updated Geo Place",
                        "description": "Updated via GeoJSON",
                        "category": "market",
                        "address": "123 Existing Street",
                        "status": "approved",
                    },
                }
            ],
        }

        importer = PlaceImporter(user=self.user, update_existing=True)
        result = importer.import_json(json.dumps(geojson_data))

        self.assertEqual(result.updated_count, 1)

        self.existing_place.refresh_from_db()
        self.assertEqual(self.existing_place.name, "Updated Geo Place")
        self.assertEqual(self.existing_place.category, "market")

    def test_import_geojson_with_metadata(self):
        """Test GeoJSON with metadata field."""
        geojson_data = {
            "type": "FeatureCollection",
            "metadata": {
                "generated_at": "2024-01-01T00:00:00",
                "count": 1,
                "source": "Test",
            },
            "features": [
                {
                    "type": "Feature",
                    "geometry": {"type": "Point", "coordinates": [127.0, 37.55]},
                    "properties": {
                        "name": "Place with Metadata",
                        "description": "Test",
                        "category": "restaurant",
                        "address": "123 Test St",
                        "status": "pending",
                    },
                }
            ],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        self.assertEqual(result.created_count, 1)

    def test_import_geojson_missing_geometry(self):
        """Test GeoJSON import with missing geometry - should fail due to required location field."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    # No geometry field - location is required in model
                    "properties": {
                        "name": "No Geometry Place",
                        "description": "Test",
                        "category": "restaurant",
                        "address": "123 Test St",
                        "status": "pending",
                    },
                }
            ],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        # Should fail because location is a required field
        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.error_count, 1)

    def test_import_geojson_nested_places_key(self):
        """Test import from JSON with 'places' key."""
        json_data = {
            "places": [
                {
                    "name": "Nested Place",
                    "description": "From nested structure",
                    "category": "restaurant",
                    "address": "123 Nested St",
                    "latitude": 37.55,
                    "longitude": 127.0,
                    "status": "pending",
                }
            ]
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(json_data))

        self.assertEqual(result.created_count, 1)

        place = HalalPlace.objects.get(name="Nested Place")
        self.assertEqual(place.address, "123 Nested St")

    def test_import_geojson_empty_features(self):
        """Test GeoJSON with empty features array."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        self.assertEqual(result.created_count, 0)
        self.assertEqual(result.error_count, 0)

    def test_import_geojson_missing_required_fields_in_properties(self):
        """Test validation of required fields in GeoJSON properties."""
        geojson_data = {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    "geometry": {"type": "Point", "coordinates": [127.0, 37.55]},
                    "properties": {
                        # Missing name and address
                        "description": "Invalid",
                        "category": "restaurant",
                        "status": "pending",
                    },
                }
            ],
        }

        importer = PlaceImporter(user=self.user)
        result = importer.import_json(json.dumps(geojson_data))

        # Validation returns early after first error (name is checked before address)
        self.assertEqual(result.error_count, 1)
        self.assertEqual(result.errors[0]["field"], "name")
